"""
Servidor HTTP y Gestor Automático de Excel para Boda Gloria & Marco
Puerto: 8089
Maneja archivos estáticos y API POST /api/rsvp para escribir directamente en confirmados_boda.xlsx
"""

import os
import json
import http.server
import socketserver
from urllib.parse import urlparse
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PORT = 8089
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE_DIR, "confirmados_boda.xlsx")

# Estilos de Gala Nupcial para el Excel
GOLD_FILL = PatternFill(start_color="C5A880", end_color="C5A880", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="8F6B3E")
ROW_FONT = Font(name="Calibri", size=10, color="2B2625")
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style='thin', color='E2CEB5'),
    right=Side(style='thin', color='E2CEB5'),
    top=Side(style='thin', color='E2CEB5'),
    bottom=Side(style='thin', color='E2CEB5')
)

COLUMNS = [
    ("N°", 6),
    ("Fecha y Hora", 18),
    ("Nombre Completo", 30),
    ("Teléfono / Celular", 18),
    ("¿Asistirá?", 12),
    ("Pases", 10),
    ("Canción Solicitada", 26),
    ("Dedicatoria / Mensaje", 35),
    ("Código Pase Digital", 18)
]

def init_excel():
    """Inicializa el archivo Excel si no existe con dos hojas y formato de gala"""
    if os.path.exists(EXCEL_PATH):
        try:
            wb = openpyxl.load_workbook(EXCEL_PATH)
            return wb
        except Exception as e:
            print(f"Error cargando Excel existente: {e}")

    wb = openpyxl.Workbook()
    
    # Hoja 1: Solo personas que SÍ confirman
    ws_si = wb.active
    ws_si.title = "Confirmados (Asistencia SÍ)"
    setup_sheet(ws_si, "LISTA OFICIAL DE INVITADOS CONFIRMADOS - BODA GLORIA & MARCO 💍")

    # Hoja 2: Todas las respuestas
    ws_all = wb.create_sheet(title="Todas las Respuestas")
    setup_sheet(ws_all, "REGISTRO TOTAL DE RESPUESTAS RSVP - BODA GLORIA & MARCO")

    wb.save(EXCEL_PATH)
    return wb

def setup_sheet(ws, title_text):
    """Aplica formato visual a una hoja de cálculo"""
    # Título en fila 1
    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = title_text
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # Encabezados en fila 3
    ws.row_dimensions[3].height = 28
    for col_idx, (col_name, col_width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=3, column=col_idx, value=col_name)
        cell.fill = GOLD_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width

def add_rsvp_to_excel(rsvp_data):
    """Agrega un registro al archivo Excel en las hojas correspondientes"""
    wb = init_excel()
    
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    attendance = str(rsvp_data.get("attendance", "SI")).upper()
    is_attending = attendance == "SI" or "SÍ" in attendance

    row_values = [
        None, # Se calcula consecutivo
        now_str,
        rsvp_data.get("name", "Invitado"),
        rsvp_data.get("phone", ""),
        "SÍ" if is_attending else "NO",
        rsvp_data.get("tickets", "2") if is_attending else "0",
        rsvp_data.get("song", ""),
        rsvp_data.get("message", ""),
        rsvp_data.get("codigoPase", "PAS-GM-CONFIRMED")
    ]

    # 1. Agregar a Todas las Respuestas
    ws_all = wb["Todas las Respuestas"]
    next_row_all = ws_all.max_row + 1
    consecutivo_all = max(1, next_row_all - 3)
    row_values[0] = consecutivo_all
    append_formatted_row(ws_all, next_row_all, row_values)

    # 2. Si SÍ asiste, agregar a la hoja principal de Confirmados
    if is_attending:
        ws_si = wb["Confirmados (Asistencia SÍ)"]
        next_row_si = ws_si.max_row + 1
        consecutivo_si = max(1, next_row_si - 3)
        row_values[0] = consecutivo_si
        append_formatted_row(ws_si, next_row_si, row_values)

    wb.save(EXCEL_PATH)
    print(f"[EXCEL ACTUALIZADO] Se guardó registro de: {rsvp_data.get('name')} (Asistencia: {row_values[4]})")
    return True

def append_formatted_row(ws, row_idx, values):
    ws.row_dimensions[row_idx].height = 22
    for col_idx, val in enumerate(values, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.font = ROW_FONT
        cell.border = THIN_BORDER
        if col_idx in (1, 2, 4, 5, 6, 9):
            cell.alignment = CENTER_ALIGN
        else:
            cell.alignment = LEFT_ALIGN

class WeddingRequestHandler(http.server.SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=BASE_DIR, **kwargs)

    def do_POST(self):
        url_parts = urlparse(self.path)
        if url_parts.path == "/api/rsvp":
            content_length = int(self.headers.get('Content-Length', 0))
            post_data = self.rfile.read(content_length)
            
            try:
                rsvp_data = json.loads(post_data.decode('utf-8'))
                add_rsvp_to_excel(rsvp_data)
                
                response = {
                    "success": True,
                    "message": "Registro guardado exitosamente en Excel",
                    "timestamp": datetime.now().isoformat()
                }
                self.send_response(200)
                self.send_header('Content-Type', 'application/json')
                self.send_header('Access-Control-Allow-Origin', '*')
                self.end_headers()
                self.wfile.write(json.dumps(response).encode('utf-8'))
            except Exception as e:
                print(f"Error procesando RSVP: {e}")
                self.send_response(500)
                self.send_header('Content-Type', 'application/json')
                self.end_headers()
                self.wfile.write(json.dumps({"success": False, "error": str(e)}).encode('utf-8'))
        else:
            self.send_error(404, "Endpoint no encontrado")

    def do_GET(self):
        url_parts = urlparse(self.path)
        if url_parts.path == "/api/confirmados":
            # Retorna resumen de confirmados
            try:
                wb = init_excel()
                ws_si = wb["Confirmados (Asistencia SÍ)"]
                count = max(0, ws_si.max_row - 3)
                self.send_response(200)
                self.send_header('Content-Type', 'application/json')
                self.send_header('Access-Control-Allow-Origin', '*')
                self.end_headers()
                self.wfile.write(json.dumps({"success": True, "confirmados_count": count}).encode('utf-8'))
            except Exception as e:
                self.send_response(500)
                self.end_headers()
        else:
            range_header = self.headers.get('Range')
            if range_header and range_header.startswith('bytes='):
                self.handle_range_request(range_header)
            else:
                super().do_GET()

    def handle_range_request(self, range_header):
        path = self.translate_path(self.path)
        if not os.path.exists(path) or os.path.isdir(path):
            super().do_GET()
            return
        
        file_size = os.path.getsize(path)
        range_val = range_header.replace('bytes=', '').strip()
        parts = range_val.split('-')
        start = int(parts[0]) if parts[0] else 0
        end = int(parts[1]) if len(parts) > 1 and parts[1] else file_size - 1
        
        if start >= file_size or end >= file_size or start > end:
            self.send_response(416)
            self.send_header('Content-Range', f'bytes */{file_size}')
            self.end_headers()
            return
        
        chunk_len = end - start + 1
        content_type = self.guess_type(path)
        
        self.send_response(206)
        self.send_header('Content-Type', content_type)
        self.send_header('Content-Range', f'bytes {start}-{end}/{file_size}')
        self.send_header('Content-Length', str(chunk_len))
        self.send_header('Accept-Ranges', 'bytes')
        self.send_header('Access-Control-Allow-Origin', '*')
        super().end_headers()
        
        try:
            with open(path, 'rb') as f:
                f.seek(start)
                bytes_left = chunk_len
                bufsize = 64 * 1024
                while bytes_left > 0:
                    read_amount = min(bytes_left, bufsize)
                    data = f.read(read_amount)
                    if not data:
                        break
                    self.wfile.write(data)
                    bytes_left -= len(data)
        except (ConnectionResetError, BrokenPipeError):
            pass

    def end_headers(self):
        self.send_header('Accept-Ranges', 'bytes')
        super().end_headers()

    def do_OPTIONS(self):
        self.send_response(200)
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'POST, GET, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        self.end_headers()

def run_server():
    init_excel()
    socketserver.ThreadingTCPServer.allow_reuse_address = True
    with socketserver.ThreadingTCPServer(("", PORT), WeddingRequestHandler) as httpd:
        print(f"Servidor de Boda activo en http://localhost:{PORT}")
        print(f"Excel vinculado: {EXCEL_PATH}")
        httpd.serve_forever()

if __name__ == "__main__":
    run_server()
